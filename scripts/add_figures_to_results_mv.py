#!/usr/bin/env python3
"""
Insert the four multi-video figures into Results_multiple_videos.docx, directly
after the abstract. Each figure gets a caption and one paragraph of explanation.

Nothing already in the document is changed. The figures are placed after the
abstract paragraph rather than appended at the end of the file.
"""
import argparse, copy, os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

ap = argparse.ArgumentParser()
ap.add_argument("--doc", default="Results_multiple_videos.docx")
ap.add_argument("--figdir", default=".")
ap.add_argument("--after", type=int, default=5, help="index of the abstract paragraph")
A = ap.parse_args()

WIDTH = 6.0

FIGS = [
    ("fig_mv1_corpus.png",
     "The corpus used in this study: 21 clips from 11 independent filming shoots.",
     "This figure shows what the corpus contains. Each bar is one clip, and the length of the "
     "bar is the number of frames taken from that clip at one frame per second. The colour "
     "gives the shoot the clip came from, and the letter at the start of each label repeats "
     "that shoot. The 21 clips together hold 460 frames. The shoots are very uneven, because "
     "shoot B supplied six clips while several shoots supplied only one. Clips from the same "
     "shoot share a room, a patient and a camera, so they are not independent of each other. "
     "This is the reason every later test in this study holds out a whole shoot rather than a "
     "single clip."),
    ("fig_mv2_reversal.png",
     "The difference between the score of the forward clip and the score of the same clip "
     "played backward, for all 21 clips.",
     "This figure shows the central result. Each bar is one clip, and its length is the "
     "difference between the score the pipeline gives to the forward clip and the score it "
     "gives to the same clip played backward. The scale is logarithmic, so each step to the "
     "right is ten times larger. The red dashed line marks the precision of the numbers the "
     "computer uses, and the green dotted line marks the size of a difference that would "
     "matter, which is about 0.01. Every bar lies far to the left of the green line, and five "
     "clips show a difference of exactly zero. The largest difference found anywhere is "
     "1.19 × 10⁻⁷, which is the size of a rounding error and not a real "
     "difference. Reversing a clip turns every action into its opposite, so this figure shows "
     "that the pipeline cannot tell an action from its opposite at all. This is not a weakness "
     "that more data would fix, because the frames of the reversed clip are the same frames in "
     "a different order, and the method never looks at the order."),
    ("fig_mv3_null_tests.png",
     "Five order-sensitive tests, each measured against chance.",
     "This figure shows what happened when the pipeline was changed so that the order of the "
     "frames does matter. Each bar is one test, and its height is the share of clips given the "
     "correct direction. The red dashed line at 0.5 is chance, because each clip can be "
     "forward or backward and a coin would be right half of the time. The thin black lines are "
     "95 per cent confidence intervals, and the text above each bar gives the raw count and "
     "the probability of seeing that result by luck alone. Every bar sits at or near the chance "
     "line, and every p value is large, which means no test did better than guessing. The "
     "confidence intervals are wide because the corpus is small, so this figure shows an "
     "absence of evidence for these simple methods rather than proof that no method can ever "
     "work."),
    ("fig_mv4_orthogonality.png",
     "Why the simple repairs failed: movement between frames is almost at a right angle to "
     "the direction in which language separates an action from its reverse.",
     "This figure explains the previous one. For each pair of neighbouring frames the study "
     "measured the direction the clip moves in the embedding space, and compared it with the "
     "direction that separates the two opposite captions. The left panel is a histogram of the "
     "angle between those two directions for every step in four clips sampled at four frames "
     "per second. An angle of 90 degrees, marked by the red dashed line, means the movement "
     "carries no information about which caption is correct. The measured angles sit almost on "
     "on top of that line: the mean over all 348 steps is 90.01 degrees, and the four "
     "clips give 89.82, 90.14, 90.21 and 89.98 degrees. The right panel gives the average projection for each clip, and the "
     "values are about 0.003, which is very close to zero and changes sign from clip to clip. "
     "In plain words, the clip does move through the embedding space as time passes, but it "
     "moves sideways with respect to the question being asked, so no simple linear read-out "
     "can recover the direction of the action."),
]

doc = Document(A.doc)
anchor = doc.paragraphs[A.after]._p

def place(el):
    """Insert el straight after the running anchor, then become the anchor."""
    global anchor
    anchor.addnext(el)
    anchor = el

h = doc.add_paragraph(style="Heading 2")
h.add_run("Figures")
place(h._p)

for i, (fn, caption, body) in enumerate(FIGS, 1):
    path = os.path.join(A.figdir, fn)
    if not os.path.isfile(path):
        raise SystemExit(f"missing figure: {path}")

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(path, width=Inches(WIDTH))
    place(p._p)

    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(6)
    r = cap.add_run(f"Figure {i}. {caption}")
    r.italic = True; r.font.size = Pt(9.5)
    place(cap._p)

    t = doc.add_paragraph(style="Normal (Web)")
    t.paragraph_format.space_after = Pt(14)
    t.add_run(body).font.size = Pt(11)
    place(t._p)

doc.save(A.doc)
print("UPDATED", A.doc, os.path.getsize(A.doc), "bytes;", len(FIGS), "figures inserted")
