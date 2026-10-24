#!/usr/bin/env python3
"""
Build the NLVL pilot results report, plain-language edition.

Differences from make_results_docx.py:
  * every sentence is short and uses ordinary words
  * each embedded figure gets a numbered caption AND a paragraph that says what
    is on each axis, what the colours mean, and what the reader should look at
  * a new section lists all 18 figure files and the frame sheets, and explains
    what each one represents
  * the claim about refusing to answer is split per model, because refusal was
    free only for the best model and cost accuracy for the weaker ones

Paths are arguments so the same file runs anywhere.
"""
import argparse, json, os, datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ap = argparse.ArgumentParser()
ap.add_argument("--root", default=os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                help="folder holding data/ and outputs/")
ap.add_argument("--figures", default=None, help="folder holding the png figures")
ap.add_argument("--out", default=None, help="output .docx path")
A = ap.parse_args()
ROOT = A.root
FIG = A.figures or f"{ROOT}/figures"
OUT = A.out or f"{ROOT}/docs/NLVL_pilot_results.docx"

NICE = {"clip": "CLIP ViT-B/32", "clip_l14": "CLIP ViT-L/14",
        "clip_l14_squash": "CLIP ViT-L/14 (squash)", "clip_l14_zoom": "CLIP ViT-L/14 (zoom)",
        "biomedclip": "BioMedCLIP ViT-B/16"}
ORDER = ["clip", "clip_l14", "clip_l14_squash", "clip_l14_zoom", "biomedclip"]

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(8); st.paragraph_format.line_spacing = 1.15
for lvl, size in ((1, 15), (2, 12.5)):
    h = doc.styles[f"Heading {lvl}"]
    h.font.name = "Calibri"; h.font.size = Pt(size); h.font.bold = True
    h.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
    h.paragraph_format.space_before = Pt(14); h.paragraph_format.space_after = Pt(6)

FIGNO = [0]


def para(t, bold=False, italic=False, size=None, align=None, after=None, indent=None):
    p = doc.add_paragraph(); r = p.add_run(t); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if align is not None: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if indent is not None: p.paragraph_format.left_indent = Inches(indent)
    return p


def bullet(t):
    p = doc.add_paragraph(t, style="List Bullet")
    p.paragraph_format.space_after = Pt(5)
    return p


def table(rows, widths=None, fs=9.5):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Light Grid Accent 1"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = ""
            pp = c.paragraphs[0]; pp.paragraph_format.space_after = Pt(2)
            rr = pp.add_run(str(v)); rr.font.size = Pt(fs); rr.bold = (i == 0)
    if widths:
        for j, w in enumerate(widths):
            for i in range(len(rows)):
                t.cell(i, j).width = w
    doc.add_paragraph()
    return t


def figure(name, title, reading, width=6.4):
    """Place a figure, then a caption, then a plain description of how to read it."""
    p = f"{FIG}/{name}"
    if not os.path.isfile(p):
        print("MISSING FIGURE", p)
        return
    FIGNO[0] += 1
    doc.add_picture(p, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(4)
    r = cap.add_run(f"Figure {FIGNO[0]}. {title}")
    r.italic = True; r.font.size = Pt(9)
    d = doc.add_paragraph(); d.paragraph_format.left_indent = Inches(0.3)
    d.paragraph_format.space_after = Pt(14)
    rr = d.add_run("How to read this figure. "); rr.bold = True; rr.font.size = Pt(9.5)
    rr2 = d.add_run(reading); rr2.font.size = Pt(9.5)


# ---------- load results ----------
R = {}
for m in ORDER:
    f = f"{ROOT}/outputs/{m}/results.json"
    if os.path.isfile(f):
        R[m] = json.load(open(f))
OPP = {r["model"]: r for r in json.load(open(f"{ROOT}/outputs/oppositional_analysis.json"))}
GT = json.load(open(f"{ROOT}/data/ground_truth.json"))
QS = {q["id"]: q for q in json.load(open(f"{ROOT}/data/queries.json"))["queries"]}
THR = json.load(open(f"{ROOT}/outputs/clip_l14/predictions.json"))["thresholds"]

# ---------- title ----------
para("Finding the Right Moment in a Medical Video from a Written Sentence",
     bold=True, size=17, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
para("Pilot Study: Results, explained in plain language", italic=True, size=12.5,
     align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
para("Mehran Piran", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para(datetime.date.today().strftime("%d %B %Y"), size=9.5, italic=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, after=14)

# ================= 1 =================
doc.add_heading("1. The Question and the Short Answer", level=1)
para(
    "This pilot study asks one question. If a person writes a sentence such as “a nurse "
    "inserts the needle into the patient’s arm”, can a computer find the exact seconds "
    "in a video where that action happens? The video is one continuous recording of a blood "
    "draw. It lasts 84 seconds. No model was trained in this study. Every model was used "
    "exactly as it was downloaded, so the results show what these models already know."
)
para(
    "The short answer is that the computer can do this only partly. It finds the general part "
    "of the video that matches the sentence. It often gets the start and the end time wrong. "
    "Most important of all, it cannot tell an action apart from the opposite of that action. "
    "This last point is the main result of the study, and Section 7 explains it."
)
para("The study found four things.", bold=True)
for n, t in enumerate([
    "The size of the image model matters more than anything else that was tested. A larger "
    "image model raised the average time overlap from 0.036 to 0.265. That is about seven "
    "times better. The small model worked no better than guessing, and no later step could "
    "fix it.",
    "Averaging the scores of neighbouring seconds gave the largest single improvement. For "
    "the large model it raised the average overlap from 0.079 to 0.265, which is 3.4 times "
    "better. For the small model it changed almost nothing. So this step only helps when the "
    "model already produces a useful signal.",
    "Letting the system say “I do not know” was free for the best model. It rejected "
    "all three easy wrong queries and one of the four difficult wrong queries, and it still "
    "found the real actions just as well as before. For the weaker models the same rule "
    "refused to answer some real actions, so it was not free for them.",
    "The model cannot tell an action from its opposite. The sentence about pulling the needle "
    "out of the arm describes something that never happens in this video. Even so, its score "
    "curve is almost the same as the curve for inserting the needle. The two curves agree with "
    "a correlation of 0.97, and they reach their highest point at the same second.",
], 1):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(7)
    p.add_run(f"Finding {n}.  ").bold = True
    p.add_run(t)

# ================= 2 =================
doc.add_heading("2. Words Used in This Report", level=1)
para(
    "This report avoids technical language where it can. A few terms are still needed, so they "
    "are defined here in simple words."
)
table([
    ("Term", "Plain meaning"),
    ("Query", "A short sentence that describes one action, written by a person."),
    ("Frozen model", "A model that is used as it was downloaded. Nothing about it is changed."),
    ("Similarity score",
     "A number that says how well one picture matches one sentence. A higher number means a "
     "better match."),
    ("Score curve",
     "The similarity score of one sentence plotted for every second of the video."),
    ("Time overlap (tIoU)",
     "How much the predicted time interval and the true time interval share. The shared "
     "seconds are divided by all the seconds that either interval covers. The value is 1 for "
     "a perfect answer and 0 when the two intervals do not touch."),
    ("R@0.3 and R@0.5",
     "The share of actions whose time overlap is above 0.3 or above 0.5."),
    ("Hard negative",
     "A sentence that uses the same words and objects as a true action, but describes "
     "something the video never shows."),
    ("Easy negative", "A sentence that describes a completely different procedure."),
], [Inches(1.5), Inches(4.7)], fs=9)

# ================= 3 =================
doc.add_heading("3. The Video, the Sentences, and the True Times", level=1)
para(
    "The video is a free clip of a blood test. It is 1280 by 720 pixels in size and it runs at "
    "23.976 frames per second. The whole clip lasts 84.42 seconds. One picture was taken every "
    "second, which gives 84 pictures. The model therefore sees 84 still pictures and never sees "
    "motion."
)
para(
    "Thirteen sentences were written by hand. Six of them describe steps that really happen in "
    "the video. Four of them are hard negatives. They reuse the same words and objects, but "
    "they describe something that never happens. Three of them are easy negatives, because they "
    "describe a different procedure such as an MRI scan or an X-ray."
)
para(
    "For each of the six real steps, the true start second and the true end second were marked "
    "by hand. The rule used to choose each boundary was written down next to it, so another "
    "person can check the same decision."
)
rows = [("Query", "Action", "True time (s)", "Length")]
for qid in ["p1", "p2", "p3", "p4", "p5", "p6"]:
    g = GT["intervals"][qid]
    rows.append((qid, QS[qid]["name"].replace("_", " "),
                 f"{g['start']} – {g['end']}", f"{g['end']-g['start']+1} s"))
table(rows, [Inches(0.8), Inches(2.3), Inches(1.6), Inches(1.1)])
para("Table 1. The six real actions and the seconds they cover. Together they cover 53 of "
     "the 84 seconds of the video.",
     italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

para("Two corrections to the first plan.", bold=True)
para(
    "All 84 pictures were looked at by eye before any model was run. That inspection changed "
    "the plan in two ways. Both changes made the test harder and more honest."
)
para(
    "The first change is about the tourniquet. The blue tourniquet is already tied around the "
    "arm in the first picture, and it is still tied in the last picture. It is never tied and "
    "never untied while the camera is running. The two sentences about tying and untying the "
    "tourniquet were therefore moved out of the real actions and into the hard negatives. They "
    "are very demanding negatives, because the object in the sentence is visible in every "
    "single second of the video. The model has to reject the action while the object is on the "
    "screen."
)
para(
    "The second change is about the number of tubes. The blood is collected into two tubes, not "
    "one. The first tube fills between 45 and 72 seconds. It is taken off at about 73 seconds. "
    "A second tube is put on and starts to fill at about 82 seconds. Only the first fill was "
    "marked as the true time for the collection sentence. The second fill is written in the "
    "ground-truth file as a known ambiguity. It was not quietly joined to the first interval, "
    "because an answer that lands in the second fill looks correct to a human eye while it "
    "counts as wrong against the marked time.",
    after=12)

# ================= 4 =================
doc.add_heading("4. How the Method Works, Step by Step", level=1)
para(
    "The method has four steps. Each step is simple on its own, and each step was measured "
    "separately so that its effect can be seen."
)
for n, (ttl, txt) in enumerate([
    ("Encode", "Every one of the 84 pictures is turned into a list of numbers by the image "
               "part of the model. Every one of the 13 sentences is turned into a list of "
               "numbers of the same length by the text part of the model."),
    ("Score", "For each sentence and each second, the two lists of numbers are compared. The "
              "result is one similarity score. This gives one score curve per sentence, with "
              "84 points on it. This first system is called E1."),
    ("Smooth", "Each point on the curve is replaced by a weighted average of itself and its "
               "two neighbours. The centre second gets half of the weight and each neighbour "
               "gets a quarter. This removes small jumps that come from single noisy pictures. "
               "This second system is called E2."),
    ("Choose an interval and decide whether to answer",
     "The highest point of the curve is found. The interval then grows to the left and to the "
     "right while the curve stays above a level set at 70 per cent of the way from the lowest "
     "to the highest value of that curve. The system also decides whether to answer at all. It "
     "answers only when the highest score passes a threshold. The threshold was fitted on the "
     "three easy negatives only, so the four hard negatives stayed an untouched test. This "
     f"third system is called E3. For the best model the threshold was {THR['tau_absolute']:.3f}."),
], 1):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(7)
    p.add_run(f"Step {n}. {ttl}.  ").bold = True
    p.add_run(txt)
para(
    "Five frozen settings were compared in this way. Two of them are general models of "
    "different size. Two of them change how the picture is prepared before the model sees it. "
    "One of them is a model trained on medical images. Table 2 reports all of them.", after=12)

# ================= 5 =================
doc.add_heading("5. Main Results", level=1)
para(
    "Table 2 gives the average time overlap for the six real actions, the share of actions "
    "above two overlap levels, and the ability to reject sentences that describe actions the "
    "video does not contain. The rejection numbers are given separately for the hard negatives "
    "and the easy negatives, because these two groups are very different in difficulty."
)
rows = [("Setting", "E1 overlap", "E2 overlap", "E3 overlap",
         "E2 R@0.3", "E2 R@0.5", "E3 reject hard", "E3 reject easy", "E3 wrong refusals")]
for m in ORDER:
    if m not in R: continue
    s = R[m]["summary"]
    rows.append((NICE[m],
                 f"{s['E1_baseline']['mean_tIoU']:.3f}",
                 f"{s['E2_temporal']['mean_tIoU']:.3f}",
                 f"{s['E3_abstention']['mean_tIoU']:.3f}",
                 f"{s['E2_temporal']['R@0.3']:.2f}",
                 f"{s['E2_temporal']['R@0.5']:.2f}",
                 f"{s['E3_abstention']['rejection_acc_hard']:.2f}",
                 f"{s['E3_abstention']['rejection_acc_easy']:.2f}",
                 f"{s['E3_abstention']['false_abstain_on_positives']} of 6"))
table(rows, fs=8)
para("Table 2. All five frozen settings and all three systems. The last column counts the "
     "real actions that the system wrongly refused to answer. The best system is the large "
     "general model with smoothing.",
     italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

figure("fig4_summary.png",
       "Accuracy and rejection ability for every setting and every system.",
       "The figure has two panels. In both panels the horizontal axis holds the three systems "
       "in order, E1 on the left, then E2, then E3. Each coloured bar is one of the five frozen "
       "settings, and the legend names them. In the left panel the height of the bar is the "
       "average time overlap, so taller is better. The clear pattern is that the bars for the "
       "large model are much taller than the rest, and that every bar grows when smoothing is "
       "added at E2. In the right panel the height of the bar is the share of wrong sentences "
       "that the system refused to answer. The bars are zero for E1 and E2, because those two "
       "systems always answer. They rise at E3, because only E3 is allowed to refuse. Reading "
       "the two panels together shows the main message of this section. For the large general "
       "model the left bar does not fall when the right bar rises, which means the system "
       "gained the ability to refuse without losing accuracy.")

para(
    "The two cropping settings were included to test a worry about the input rather than about "
    "the model. The standard picture pipeline makes the shorter side of the picture 224 pixels "
    "long and then cuts a square out of the centre. On a picture that is 1280 by 720 pixels "
    "this keeps only 56 per cent of the width. The left and the right part of the picture are "
    "thrown away before the model sees anything. The squash setting avoids the cut by pressing "
    "the whole picture into a square. The zoom setting cuts a tighter box around the place "
    "where the procedure happens."
)
para(
    "Squashing made the result much worse, from 0.265 down to 0.085. This suggests that keeping "
    "the shape of the picture similar to the pictures the model was trained on matters more "
    "than keeping all of the content. Zooming almost doubled the first system, from 0.079 to "
    "0.151, but it ended below the standard pipeline once smoothing was added, at 0.201 against "
    "0.265. Zooming also caused two wrong refusals. The way the picture is prepared is "
    "therefore a real factor, but it is a smaller factor than the size of the model."
)
para(
    "The medical model finished well behind the large general model, at 0.098 against 0.265. "
    "This comparison needs care. The medical model uses a smaller image encoder, so part of the "
    "gap is only a difference in size. Against the general model of a similar size it is "
    "clearly better, at 0.098 against 0.036. The fair conclusion is that medical training does "
    "help at a fixed model size, but it does not make up for a smaller encoder. The pictures "
    "the medical model learned from are figures printed in scientific papers, and those are a "
    "poor match for video recorded in a clinic.", after=12)

# ================= 6 =================
doc.add_heading("6. Results for Each Action", level=1)
per = R["clip_l14"]["per_query"]["E2_temporal"]
rows = [("Query", "Action", "True time", "Predicted time", "Overlap")]
for qid in ["p1", "p2", "p3", "p4", "p5", "p6"]:
    r = per[qid]
    rows.append((qid, r["name"].replace("_", " "),
                 f"{r['gt'][0]} – {r['gt'][1]} s",
                 f"{r['pred'][0]:.0f} – {r['pred'][1]:.0f} s",
                 f"{r['tIoU']:.3f}"))
table(rows, [Inches(0.7), Inches(1.9), Inches(1.4), Inches(1.4), Inches(0.8)])
para("Table 3. Results for each action, for the best system. Cleaning the arm is found well. "
     "Preparing the needle and changing the tube fail completely.",
     italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

para(
    "The average hides two very different kinds of behaviour. Cleaning the arm is found well, "
    "with an overlap of 0.714. Collecting the blood and attaching the tube are found in roughly "
    "the right place, but the predicted interval is too short. Two actions fail completely with "
    "an overlap of zero. Both of them failed in the same way. Their highest score falls in the "
    "middle part of the video, between 25 and 36 seconds, where the hands work close to the arm "
    "and the pictures all look alike. Changing the tube happens at 73 to 77 seconds, yet the "
    "system answered 31 to 41 seconds. The model matched the objects in the sentence and ignored "
    "when the action takes place."
)

figure("clip_l14_fig1_positive_curves.png",
       "Score curves for the six real actions, for the best system.",
       "There is one panel for each of the six sentences, and the name of the action is printed "
       "on the left of its panel. In every panel the horizontal axis is time in seconds, from 0 "
       "to 84, and the vertical axis is the similarity score after smoothing. The dark line is "
       "the score curve of that sentence. The green band marks the true seconds of the action. "
       "The orange band marks the seconds the system predicted. The number in the top right "
       "corner of each panel is the time overlap for that action. A good result looks like a "
       "green band and an orange band sitting on top of each other, which happens in the first "
       "panel. A failure looks like an orange band far away from the green band, which happens "
       "in the second panel and in the last panel. Note that the vertical range of each panel "
       "is small. The curves move by only a few hundredths, so the useful signal is weak even "
       "when it points to the right place.")

figure("clip_l14_fig3_smoothing_effect.png",
       "What averaging over neighbouring seconds does to one curve.",
       "The horizontal axis is time in seconds and the vertical axis is the similarity score. "
       "The grey line is the raw curve, where each second is scored on its own. The dark line is "
       "the same curve after each second was averaged with its two neighbours. The green band is "
       "the true time of the action, which here is cleaning the arm. The grey line jumps up and "
       "down from one second to the next, and several of those jumps are as large as the real "
       "peak. The dark line keeps the peak but removes most of the jumps. This is the whole "
       "reason the average overlap rose from 0.079 to 0.265. The step adds no new information "
       "about the video. It only stops single noisy pictures from winning the argument.")

# ================= 7 =================
doc.add_heading("7. The Main Result: an Action and Its Opposite", level=1)
para(
    "The most useful result of this pilot is not a score. It is a failure of a very particular "
    "kind, and it is easy to state."
)
para(
    "One of the hard negatives says that the nurse pulls the needle out of the arm. This never "
    "happens in the video. The needle is still in the arm in the last picture. Even so, the "
    "score curve of that sentence is almost the same as the score curve of the sentence that "
    "says the nurse inserts the needle. For the best model the two curves agree with a "
    "correlation of 0.969, and both reach their highest value at the same second."
)
para(
    "This is not a general inability to tell sentences apart. The same model separates the "
    "needle sentence from a sentence about an MRI scanner almost perfectly. Their correlation is "
    "only 0.062, and their highest points are 39 seconds apart. So the model knows very well "
    "which procedure it is looking at. What it cannot do is tell which way an action runs when "
    "the objects in the sentence stay the same."
)
para(
    "The reason is simple. The model scores each second on its own, and a single picture of a "
    "needle in an arm does not say whether the needle is going in or coming out. Direction lives "
    "in the change between pictures, and this method never looks at the change."
)
rows = [("Setting", "Agreement with the opposite action",
         "Agreement with an unrelated action", "Distance between the two peaks")]
for m in ORDER:
    if m not in OPP: continue
    o = OPP[m]
    rows.append((NICE[m], f"{o['r_opposite']:.3f}", f"{o['r_unrelated']:.3f}",
                 f"{o['peak_dist_opposite']} s"))
table(rows, [Inches(1.7), Inches(1.7), Inches(1.6), Inches(1.5)])
para("Table 4. The sentence about inserting the needle compared with the sentence about "
     "pulling it out, and then compared with a sentence about an MRI scanner. A value near 1 "
     "means the two curves move together. A value near 0 means they move independently.",
     italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

figure("fig2b_oppositional_with_reference.png",
       "The true sentence, its opposite, and an unrelated sentence, on one set of axes.",
       "The horizontal axis is time in seconds and the vertical axis is the similarity score "
       "after smoothing. The green band marks the seconds where the needle really is inserted. "
       "The solid dark line is the sentence about inserting the needle. The red dashed line is "
       "the sentence about pulling the needle out, which never happens in this video. The grey "
       "dotted line is the sentence about an MRI scanner, which is included as a reference for "
       "what a sentence that does not belong looks like. The point of the figure is the distance "
       "between the lines. The red dashed line follows the solid dark line almost exactly, rises "
       "and falls with it, and peaks at the same second. The grey dotted line sits far below and "
       "moves on its own. So the model places the opposite of a true action in the same place as "
       "the action itself, while it keeps a different procedure clearly apart.")

figure("biomedclip_fig2_oppositional_pair.png",
       "The same pair of sentences, scored by the medical model.",
       "The axes are the same as in the previous figure. The horizontal axis is time in seconds "
       "and the vertical axis is the similarity score. The dark line is the sentence about "
       "inserting the needle, the red dashed line is the sentence about pulling it out, and the "
       "green band is the true time of the insertion. This figure is shown because the medical "
       "model behaves differently from all the others. Its two lines still rise and fall "
       "together, but they come apart in the middle of the video, and their highest points are "
       "eleven seconds apart instead of zero. This is the only model tested that separates the "
       "two directions at all. It is also the weakest model at finding the right time. The "
       "result comes from one pair of sentences on one video, so it is a lead worth testing "
       "properly and not an established finding.")

figure("fig5_oppositional_separation.png",
       "The same comparison for all five settings.",
       "The horizontal axis holds the five frozen settings. For each setting there are two bars, "
       "and the height of a bar is the agreement between two score curves. A tall bar means the "
       "two curves move together. The dark red bar on the left of each pair compares the true "
       "sentence with its opposite. The light blue bar on the right of each pair compares the "
       "true sentence with the unrelated MRI sentence. The number printed above each bar is its "
       "exact value. The dark red bars are tall everywhere, between 0.77 and 0.98. The light "
       "blue bars are short everywhere, between minus 0.15 and 0.27. The gap between the two "
       "bars is the problem described in this section, and it does not close when the model gets "
       "larger. The large general model has the widest gap of all the general settings. Only the "
       "medical model on the right has a visibly shorter dark red bar.")

# ================= 8 =================
doc.add_heading("8. What Each Figure File Represents", level=1)
para(
    "Eighteen figures were produced and all of them are kept, because they let any single claim "
    "in this report be checked against another setting. The figures follow a fixed naming rule. "
    "The name starts with the setting, and then gives the figure type. The three figure types "
    "are the same for every setting, so the same picture can be compared across models. The "
    "six figures above are the ones needed to follow the argument. This section explains what "
    "every file represents."
)
rows = [("File name pattern", "What it represents")]
for a, b in [
    ("<setting>_fig1_positive_curves.png",
     "One panel for each of the six real actions. Time in seconds runs along the bottom and the "
     "smoothed similarity score runs up the side. The green band is the true time, the orange "
     "band is the predicted time, and the number in the corner is the overlap. Use it to see "
     "which actions a setting finds and which it misses."),
    ("<setting>_fig2_oppositional_pair.png",
     "The sentence about inserting the needle and the sentence about pulling it out, drawn on "
     "one set of axes with the true insertion time as a green band. Use it to see whether a "
     "setting can tell the two directions apart."),
    ("<setting>_fig3_smoothing_effect.png",
     "The raw curve in grey and the smoothed curve in dark blue for the cleaning action. Use it "
     "to see how noisy a setting is from one second to the next."),
    ("fig4_summary.png",
     "Two panels of bars. The left panel is the average time overlap and the right panel is the "
     "share of wrong sentences that were refused, both for the three systems and all five "
     "settings. This is the one figure that compares everything at once."),
    ("fig2b_oppositional_with_reference.png",
     "The best model only. The true sentence, its opposite, and an unrelated MRI sentence as a "
     "reference, on one set of axes. The reference line is what makes this figure stronger than "
     "the plain oppositional pair."),
    ("fig5_oppositional_separation.png",
     "Bars comparing agreement with the opposite action against agreement with an unrelated "
     "action, for all five settings. It shows the main failure is present in every setting."),
    ("annotation_evidence/sheet_01 to sheet_07",
     "Seven contact sheets. Each sheet shows twelve seconds of the video as twelve small "
     "pictures with the second printed on each one. Together the seven sheets show all 84 "
     "seconds. These are the pictures the true times were written from, so they are the evidence "
     "for Table 1."),
    ("annotation_evidence/A to H",
     "Eight single pictures chosen to settle one question each. They show the tourniquet already "
     "tied, the insertion, the tube being attached, the first blood, the tube being removed, the "
     "arm at the end, the second hand entering the frame, and the final second of the video. "
     "Picture A and picture H together prove that the tourniquet is never tied or untied on "
     "camera."),
]:
    rows.append((a, b))
table(rows, [Inches(2.1), Inches(4.3)], fs=8.5)
para("Table 5. Every figure file and what it represents. The word setting stands for one of "
     "clip, clip_l14, clip_l14_squash, clip_l14_zoom or biomedclip.",
     italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

# ================= 9 =================
doc.add_heading("9. What This Study Cannot Tell Us", level=1)
for t in [
    "The study uses one video and six marked actions. Every number therefore carries a large "
    "uncertainty, and no claim about statistical significance is made anywhere in this report. "
    "The numbers show the size and the direction of an effect and nothing more.",
    "The true times were marked by one person from pictures taken once per second. A boundary "
    "can be wrong by about one second. The shortest action is inserting the needle, which lasts "
    "four seconds, so a one second error changes its overlap a great deal. This is the part of "
    "the study that most needs a second pair of eyes.",
    "The threshold for refusing to answer was fitted on three easy negatives. This keeps the "
    "four hard negatives honest, but three sentences are very little data to fit anything on.",
    "The second part of the refusal rule, which compares the best score against the next best "
    "score, carried almost no information. Its fitted value was negative for every setting, "
    "which means the simple threshold on the score was doing all of the work. A better measure "
    "of confidence is needed.",
    "The video lasts 84 seconds. That is far too short to say anything about methods built for "
    "long videos, including methods that have to decide what to keep in memory.",
    "The model never sees motion. It sees 84 separate pictures. Any conclusion in this report "
    "is a conclusion about still pictures and the sentences that match them.",
]:
    bullet(t)

# ================= 10 =================
doc.add_heading("10. What Should Be Tried Next", level=1)
para(
    "The clearest direction follows from Section 7. A frozen model of this kind matches the "
    "objects named in a sentence, and it is nearly blind to the direction of the action between "
    "them. Any method that wants to find fine steps of a procedure has to bring that missing "
    "information from somewhere else. Three sources are possible, and each one can be tested on "
    "this same small benchmark."
)
for n, t in enumerate([
    "Motion. The direction of an action is visible in the change between two pictures even when "
    "it is invisible in either picture on its own. This means either optical flow or a model "
    "trained on video rather than on still pictures.",
    "Order. The steps of a procedure happen in a fixed sequence. A model that knows the sequence "
    "can rule out an insertion that would have to come after a removal.",
    "Training against opposites. If the text part of the model is trained on pairs of sentences "
    "that differ in one word, it is forced to separate them.",
], 1):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(7)
    p.add_run(f"Source {n}.  ").bold = True
    p.add_run(t)
para(
    "A second and more practical direction is the measure of confidence. The ability to refuse to "
    "answer is what would make such a system usable in a clinic, and in this pilot it came for "
    "free for the best model. However, the rule rejected only one of the four hard negatives, "
    "and the hard negatives are exactly the cases that matter. A measure built on the shape of "
    "the score curve, rather than on its height, is the obvious next thing to try."
)
para(
    "The third direction is length. The present video is short enough that every sentence can be "
    "compared against every second directly. On a procedure that lasts tens of minutes this "
    "becomes expensive, and the question of which parts of the video to keep becomes a real one.",
    after=12)

# ================= 11 =================
doc.add_heading("11. Files", level=1)
para("Everything is kept in the project repository. The scripts are also copied "
     "to the Desktop folder NLVL_scripts.", size=10.5)
rows = [("Path", "Contents")]
for a, b in [
    ("data/", "the video, the 84 pictures, frames.csv, queries.json, ground_truth.json, source.txt"),
    ("scripts/01_encode.py", "turns pictures and sentences into numbers with a frozen model"),
    ("scripts/02_localise.py", "scores, smoothing, choice of interval, and the refusal rule"),
    ("scripts/03_evaluate.py", "time overlap, recall, and rejection accuracy"),
    ("scripts/04_figures.py", "the per-action curves, the smoothing figure, and the summary bars"),
    ("scripts/05_oppositional.py", "the analysis of an action against its opposite, Section 7"),
    ("scripts/run_all.sh", "repeats the whole study with one command"),
    ("scripts/make_results_docx_v2.py", "builds this report"),
    ("outputs/<setting>/", "numbers, scores, predictions and results for each setting"),
    ("figures/", "all 18 figures and the contact sheets used to mark the true times"),
]:
    rows.append((a, b))
table(rows, [Inches(2.1), Inches(4.3)], fs=9)

os.makedirs(os.path.dirname(os.path.abspath(OUT)), exist_ok=True)
doc.save(OUT)
print("WROTE", OUT, os.path.getsize(OUT), "bytes;", FIGNO[0], "figures embedded")
