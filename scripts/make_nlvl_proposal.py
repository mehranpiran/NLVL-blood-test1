#!/usr/bin/env python3
"""Build the NLVL pilot-project proposal docx."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import datetime, os

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "Results", "NLVL_pilot_proposal.docx")

doc = Document()

# ---- base style ----
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(8)
st.paragraph_format.line_spacing = 1.15

for lvl, size in ((1, 15), (2, 12.5)):
    h = doc.styles[f"Heading {lvl}"]
    h.font.name = "Calibri"
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)

def para(text, bold=False, italic=False, size=None, align=None, space_after=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p

def numbered(n, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(f"Step {n}. {title}  ")
    r.bold = True
    p.add_run(body)
    return p

def bullet(text):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    return p

# ---- title block ----
para("Fine-Grained Language-Guided Temporal Localisation in a Medical Procedure",
     bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
para("A One-Day Pilot Study Using Frozen Vision-Language Models",
     italic=True, size=12.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
para("Mehran Piran", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para(f"Project proposal | {datetime.date.today().strftime('%d %B %Y')}",
     size=9.5, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

# ---- 1. Summary ----
para(
    "Revision note. This proposal was written before the video was inspected frame by frame. That "
    "inspection changed two things. The tourniquet is already tied in the first frame and is never "
    "applied or removed on camera, so the two tourniquet queries became hard negatives instead of "
    "positive queries. The blood is also collected into two tubes rather than one. Both changes are "
    "recorded in the results report, which supersedes this document where they disagree.",
    italic=True, size=9.5, space_after=12)

doc.add_heading("1. Summary of the Project", level=1)
para(
    "This document describes a small pilot project in natural language video localisation (NLVL). "
    "In NLVL, the computer receives a video and a sentence written in normal language. The computer "
    "must then return the start time and the end time of the part of the video that the sentence "
    "describes. The task is useful because people usually remember what happened in a video, but they "
    "do not remember when it happened."
)
para(
    "The pilot applies this task to a medical procedure. The chosen procedure is a blood draw "
    "(phlebotomy), recorded as one continuous video. A medical procedure is a hard and interesting "
    "case because all of its steps look very similar. The room, the clothing, the lighting, the arm "
    "and the camera angle stay almost the same from the first second to the last second. The only "
    "thing that changes is the fine detail of the action, for example whether the needle is being "
    "inserted or being removed. Therefore the model cannot find the correct moment by recognising the "
    "scene. It must recognise the action itself."
)
para(
    "The project does not train any new neural network. It uses a vision-language model that is "
    "already trained, and it keeps the weights frozen. The contribution of the pilot is not a new "
    "model but a careful measurement: it shows how well a frozen model performs on this hard case, "
    "and it shows how much two simple and cheap additions improve the result. The two additions are "
    "temporal consistency (smoothing the scores over time) and calibrated abstention (allowing the "
    "system to say “this action is not in the video”)."
)

doc.add_heading("1.1 Research Question", level=2)
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.3)
p.paragraph_format.right_indent = Inches(0.3)
r = p.add_run(
    "Can a frozen vision-language model localise the fine-grained steps of a visually homogeneous "
    "medical procedure, and can simple temporal modelling and calibrated abstention improve its "
    "reliability?"
)
r.italic = True
r.bold = True
para(
    "The question has three parts. The first part asks whether the model can find the right moment at "
    "all when the scene gives it no help. The second part asks whether adding a small amount of "
    "temporal reasoning fixes the mistakes that come from judging each frame alone. The third part "
    "asks whether the system can recognise that an action was never performed, instead of always "
    "returning some interval. The third part matters most in clinical work, because a confident but "
    "wrong answer is worse than an honest refusal."
)

doc.add_heading("1.2 Aim of the Project, Step by Step", level=2)
para("The aim of the project can be written as five smaller aims that follow each other in order:")
numbered(1, "Build the task.", "Create a small but honest benchmark for fine-grained temporal localisation inside one continuous medical video, including manual ground-truth time intervals for every query.")
numbered(2, "Measure the baseline.", "Measure how accurately a frozen vision-language model locates each step of the procedure when every frame is scored independently.")
numbered(3, "Add temporal consistency.", "Add a simple smoothing step over the time axis, and measure how much the accuracy changes. Actions last for several seconds, so neighbouring frames should agree with each other.")
numbered(4, "Add calibrated abstention.", "Give the system the ability to reject a query. Measure how often it correctly refuses to answer when the described action is not present in the video, and check that this ability does not damage its accuracy on the real actions.")
numbered(5, "Report the limits.", "State clearly where the frozen model fails, which failures the simple additions repair, and which failures would require training or a stronger temporal model. This final step turns the pilot into a research direction rather than a demonstration.")

# ---- 2. Data ----
doc.add_heading("2. Data: What Type and How Much", level=1)
para(
    "The data requirement for this pilot is deliberately very small, because the project trains "
    "nothing. The only data needed is one continuous video of a real medical procedure, together with "
    "a small set of text queries and a small set of hand-made ground-truth time intervals. The video "
    "is a publicly available, freely licensed clip of a blood examination that lasts about one minute "
    "and twenty-four seconds and has a resolution of 1280 by 720 pixels. It is a single unbroken recording. This "
    "last property is essential: if the video were made by joining several short clips together, each "
    "clip would have a different background, and the model could find the correct moment simply by "
    "noticing the change of scene. A single continuous recording removes that shortcut and forces the "
    "model to judge the action."
)
para(
    "From this one video, the frames are sampled at one frame per second, which gives about eighty-four "
    "frames in total. Each frame is one data point for the model, so the effective dataset size is "
    "roughly eighty-four images. Around thirteen text queries are written by hand. Six of them are "
    "positive queries that describe fine-grained steps actually visible in the video, namely applying "
    "the tourniquet, cleaning the arm, preparing the needle, inserting the needle, collecting the "
    "blood, and finishing the procedure. Four of them are oppositional hard negatives. A hard negative "
    "uses almost exactly the same words as a positive query but describes the opposite action, for "
    "example “the nurse removes the needle from the arm” against “the nurse inserts the needle into the "
    "arm”. These queries test whether the model understands the direction of the action or only "
    "recognises the objects in the sentence. The remaining three are easy negatives that describe "
    "procedures that never appear in this video, such as an MRI scan, a blood pressure measurement and "
    "an X-ray. For every positive query, the true start and end times are annotated manually by "
    "watching the video. This annotation takes a short time but it is the part of the project that "
    "makes the numbers meaningful, so it must be done carefully and the rules used for marking the "
    "boundaries must be written down."
)
para(
    "In summary, the project needs one video of about eighty-four seconds, about eighty-four sampled "
    "frames, about thirteen text queries, and about six manually annotated ground-truth intervals. No "
    "patient data and no private clinical recordings are used, so there is no ethical approval problem "
    "and the whole pilot stays reproducible by anyone."
)

# ---- 3. Models ----
doc.add_heading("3. Pretrained Models: Which Ones and Why", level=1)
para(
    "The project uses pretrained models only, and it never updates their weights. The main model is "
    "CLIP, a vision-language model trained on a very large number of image and caption pairs collected "
    "from the internet. CLIP contains two encoders. The image encoder turns a video frame into a "
    "vector, and the text encoder turns a sentence into a vector in the same space. Because both "
    "vectors live in the same space, their similarity can be measured directly with a cosine "
    "similarity. This single property is what makes the whole pilot possible without any training: "
    "the query sentence is encoded once, every frame is encoded once, and the similarity between them "
    "becomes a score over time. CLIP is therefore used for two separate jobs, namely to describe the "
    "visual content of each frame and to describe the meaning of each query, and the comparison "
    "between the two outputs is the localisation signal itself."
)
para(
    "A second pretrained model is used as an optional comparison. BioMedCLIP has the same structure as "
    "CLIP but it was trained on about fifteen million figure and caption pairs taken from biomedical "
    "scientific articles. It is included in order to answer a specific and honest question: does a "
    "model trained on medical text and medical images actually help on real clinical video? The "
    "expected answer is that it may not help, because the images it learned from are scientific "
    "figures, microscopy pictures and radiology scans, not video frames of a nurse in a clinic. If "
    "BioMedCLIP performs worse than general CLIP, that result is still a valuable finding, because it "
    "shows that the domain of the pretraining images matters more than the domain of the pretraining "
    "words. No other pretrained network is required. In particular, no object detector, no action "
    "recognition network and no large language model is needed for the first version of this pilot."
)

# ---- 4. Method steps ----
doc.add_heading("4. What Needs To Be Done: The Full Procedure", level=1)
para("The following numbered steps describe everything that must be done, in the order in which it must be done, to reach the final results.")
steps = [
 ("Prepare the project folder.",
  "Create the project directory and its subfolders for data, scripts, outputs, figures and documents, so that every file produced by the project has a fixed and predictable place."),
 ("Obtain the video.",
  "Download the single continuous phlebotomy video at the resolution the source provides, which is 1280 by 720 pixels, and store it in the data folder together with a short text file that records its source, its licence and its exact duration."),
 ("Sample the frames.",
  "Extract one frame per second from the video and save the frames as numbered image files. Record the exact timestamp of every frame, because the timestamps are what later turn frame indices back into seconds."),
 ("Write the queries.",
  "Write the thirteen text queries in a single structured file. Mark each query with its type, which is either positive, oppositional hard negative, or easy negative. Keeping them in one file makes the experiment easy to extend later."),
 ("Annotate the ground truth.",
  "Watch the video and write down the true start time and end time of each of the six positive actions. Write down the rule that was used to decide the boundary of an action, for example that the insertion begins at the first frame in which the needle touches the skin."),
 ("Encode the frames.",
  "Pass every sampled frame through the frozen CLIP image encoder and save the resulting vectors to disk. This step is done once, and all later experiments reuse the saved vectors, which keeps the whole pipeline fast."),
 ("Encode the queries.",
  "Pass every query through the frozen CLIP text encoder and save the resulting vectors in the same way."),
 ("Compute the similarity curve.",
  "For each query, compute the cosine similarity between its text vector and the vector of every frame. The result is one score for every second of the video, which can be drawn as a curve over time."),
 ("Produce the baseline prediction.",
  "For each query, take the frames whose score is highest and convert them into a single continuous time interval. This is the baseline system, in which every frame is judged completely independently of its neighbours."),
 ("Add temporal consistency.",
  "Smooth the score curve along the time axis before choosing the interval, using a simple weighted average in which each frame receives one half of its own score, one quarter of the score of the previous frame and one quarter of the score of the next frame. This step uses the fact that a real action lasts for several seconds, so it removes isolated frames that received a high score by accident."),
 ("Add calibrated abstention.",
  "Define two conditions that must both be satisfied before the system is allowed to answer. The first condition is that the highest smoothed score must be above an absolute threshold. The second condition is that the gap between the best query and the second best query must be above a margin threshold. If either condition fails, the system returns no interval and reports that the action was not found. Choose both thresholds using the easy negatives only, so that the thresholds are not tuned on the data used for the final measurement."),
 ("Evaluate every version.",
  "Compute the evaluation metrics for the baseline, for the baseline with temporal consistency, and for the full system with temporal consistency and abstention. The metrics are the mean and median temporal intersection over union, which measures how much the predicted interval overlaps the true interval; the recall at an overlap of 0.3 and at an overlap of 0.5, which measure how many queries were localised acceptably and well; and the rejection accuracy, which measures how often the system correctly refused to answer a negative query."),
 ("Repeat with BioMedCLIP.",
  "Run the identical pipeline with BioMedCLIP in place of CLIP and compare the two sets of numbers. Report the comparison honestly in either direction."),
 ("Make the figures.",
  "Draw the score curve over time for a small number of informative queries, and mark the true interval and the predicted interval on the same axis. Draw one figure that shows a positive query and its oppositional hard negative together, because this single figure demonstrates the central difficulty of the task better than any table."),
 ("Write the results.",
  "Collect the metrics into one small table, state which additions helped and by how much, and describe the cases in which the system still failed. Explain each failure in terms of what the model could and could not see."),
 ("State the next extension.",
  "Close the report by naming the natural next step, which is to apply a memory mechanism so that the same method can work on long untrimmed procedural videos rather than on a short clip. The present video lasts only eighty-four seconds, which is too short to demonstrate a memory mechanism fairly, so this extension is proposed and not claimed."),
]
for i, (t, b) in enumerate(steps, 1):
    numbered(i, t, b)

# ---- 5. Experiments table ----
doc.add_heading("5. Planned Experiments and Expected Outputs", level=1)
rows = [
    ("Experiment", "What is added", "What it tests"),
    ("E1 – Baseline", "Frozen CLIP, frame-by-frame scoring", "Whether a frozen model can localise fine-grained steps at all"),
    ("E2 – Temporal", "E1 plus temporal smoothing", "Whether simple temporal modelling removes isolated false peaks"),
    ("E3 – Abstention", "E2 plus absolute and margin thresholds", "Whether the system can refuse unseen actions without losing accuracy"),
    ("E4 – Domain (optional)", "E3 with BioMedCLIP instead of CLIP", "Whether biomedical pretraining transfers to clinical video"),
]
tbl = doc.add_table(rows=len(rows), cols=3)
tbl.style = "Light Grid Accent 1"
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, row in enumerate(rows):
    for j, val in enumerate(row):
        c = tbl.cell(i, j)
        c.text = ""
        pp = c.paragraphs[0]
        pp.paragraph_format.space_after = Pt(3)
        rr = pp.add_run(val)
        rr.font.size = Pt(10)
        rr.bold = (i == 0)
for j, w in enumerate((Inches(1.6), Inches(2.1), Inches(2.6))):
    for i in range(len(rows)):
        tbl.cell(i, j).width = w

doc.add_paragraph()
para(
    "The final deliverables are one results table, two or three figures, the saved embeddings and "
    "scores, and the scripts that reproduce everything from the raw video in a single command.",
    size=10.5, italic=True,
)

# ---- 6. Why relevant ----
doc.add_heading("6. Relevance of the Design", level=1)
para(
    "Each design choice in this pilot connects to an open problem in language-guided video "
    "understanding. The use of one continuous and visually homogeneous video removes the scene "
    "shortcut and makes the task genuinely fine-grained. The temporal smoothing step is the simplest "
    "possible form of temporal consistency, which is known to be important when supervision is weak or "
    "absent. The oppositional hard negatives test whether the model understands the direction and the "
    "structure of an action rather than the presence of its objects, which is exactly the weakness of "
    "contrastive image and text models. The abstention mechanism addresses reliability, which is the "
    "requirement that separates a clinical system from a demonstration. Finally, the medical setting "
    "is chosen because it is where a wrong localisation has a real cost, and where an honest refusal "
    "is a useful output rather than a failure."
)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print("WROTE", OUT, os.path.getsize(OUT), "bytes")
