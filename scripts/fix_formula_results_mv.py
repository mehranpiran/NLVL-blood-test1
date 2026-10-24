#!/usr/bin/env python3
"""
Correct the typed-out LaTeX in Results_multiple_videos.docx, in place.

The abstract contains the raw LaTeX source of the value, spread over five
runs. It is replaced with 1.19 x 10 followed by a real superscript minus seven,
so the value reads as a number in Word instead of as markup. The stray repeated
file name at the end of the abstract is removed at the same time.
"""
import argparse
from docx import Document

ap = argparse.ArgumentParser()
ap.add_argument("--doc", default="Results_multiple_videos.docx")
A = ap.parse_args()

doc = Document(A.doc)
p = doc.paragraphs[5]
runs = p.runs

# run 2 ends with the opening of the formula
assert runs[2].text.endswith("\\(1.19\\times10"), runs[2].text[-40:]
runs[2].text = runs[2].text[:-len("\\(1.19\\times10")] + "1.19 × 10"

# runs 3, 4, 5 hold  ^{  -  7}\   : collapse them into one superscript
runs[3].text = "−7"
runs[3].font.superscript = True
runs[4].text = ""
runs[5].text = ""

# run 6 opens with the closing bracket of the formula
assert runs[6].text.startswith(")"), runs[6].text[:20]
runs[6].text = runs[6].text[1:]

# stray duplicated file name pasted at the end of the abstract
for i in (7, 8, 9):
    if "NLVL_multivideo_timereversal" in runs[i].text or runs[i].text == " ":
        runs[i].text = ""
runs[10].text = ""          # the full stop that belonged to that stray text
runs[6].text = runs[6].text.rstrip() + " "

doc.save(A.doc)
print("FIXED", A.doc)
print(doc.paragraphs[5].text[-420:])
